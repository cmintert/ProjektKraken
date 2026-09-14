"""One longform document snapshot rendered to Markdown, DOCX or PDF."""

from __future__ import annotations

import re
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import Element, SubElement

import markdown  # type: ignore[import-untyped]  # Markdown ships no type information.
from PySide6.QtCore import QMarginsF, QUrl
from PySide6.QtGui import QPageLayout, QPageSize, QPdfWriter, QTextDocument


class _TreeParser(HTMLParser):
    """Normalize Markdown's HTML into a small document tree for Word rendering."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root = Element("document")
        self.stack = [self.root]

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        node = SubElement(self.stack[-1], tag, {k: v or "" for k, v in attrs})
        if tag not in {"img", "br", "hr", "meta", "input"}:
            self.stack.append(node)

    def handle_endtag(self, tag: str) -> None:
        for index in range(len(self.stack) - 1, 0, -1):
            if self.stack[index].tag == tag:
                self.stack = self.stack[:index]
                return

    def handle_data(self, data: str) -> None:
        parent = self.stack[-1]
        if len(parent):
            parent[-1].tail = (parent[-1].tail or "") + data
        else:
            parent.text = (parent.text or "") + data


def document_snapshot(
    sequence: list[dict[str, Any]], options: dict[str, Any], world_path: str
) -> dict[str, Any]:
    """Capture ordered narrative, readable wiki links, and contained local images."""
    title = str(options.get("title") or "Longform document")
    names = {item["name"]: item["id"] for item in sequence}
    lines = [f"# {title}", ""]
    warnings: list[str] = []
    if options.get("contents"):
        lines += ["## Contents", ""]
        lines.extend(
            f"- [{item['meta'].get('title_override') or item['name']}]"
            f"(#item-{item['id']})"
            for item in sequence
        )
        lines.append("")
    for item in sequence:
        marker = f"<!-- PK-LONGFORM id={item['id']} table={item.get('table', '')} -->"
        lines.append(marker)
        lines.append(f'<a id="item-{item["id"]}"></a>')
        lines.append(
            "#" * item["heading_level"]
            + " "
            + (item["meta"].get("title_override") or item["name"])
        )
        lines.append("")
        content = item.get("content", "")

        def wiki_link(match: re.Match[str]) -> str:
            target, _, label = match.group(1).partition("|")
            label = label or target
            return f"[{label}](#item-{names[target]})" if target in names else label

        content = re.sub(r"(?<!!)\[\[([^\]]+)\]\]", wiki_link, content)
        lines += [content, ""]
    text = "\n".join(lines)
    rendered = markdown.markdown(
        text, extensions=["tables", "fenced_code", "sane_lists"]
    )
    tree = _TreeParser()
    tree.feed(rendered)
    root = Path(world_path).resolve() if world_path else None
    for node in tree.root.iter("img"):
        source = node.get("src", "")
        path = (root / source).resolve() if root else Path(source).resolve()
        if not options.get("images", True):
            node.tag = "span"
            node.text = node.get("alt", "")
            node.attrib.clear()
        elif root is None or not path.is_relative_to(root) or not path.is_file():
            warnings.append(f"Image omitted: {source} (missing or outside the world).")
            node.tag = "span"
            node.text = f"[Image: {node.get('alt') or source}]"
            node.attrib.clear()
        else:
            node.set("src", QUrl.fromLocalFile(str(path)).toString())
            node.set("width", "520")
    supported = {
        "document",
        "p",
        "a",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "ul",
        "ol",
        "li",
        "strong",
        "em",
        "b",
        "i",
        "code",
        "pre",
        "blockquote",
        "table",
        "thead",
        "tbody",
        "tr",
        "th",
        "td",
        "img",
        "span",
        "br",
        "hr",
    }
    for node in tree.root.iter():
        if node.tag not in supported:
            warnings.append(f"Formatting simplified: {node.tag}")
            node.tag = "span"
            node.attrib.clear()
    from xml.etree.ElementTree import tostring

    body = "".join(
        tostring(node, encoding="unicode", method="html") for node in tree.root
    )
    return {
        "markdown": text,
        "html": body,
        "title": title,
        "page_size": options.get("page_size", "A4"),
        "warnings": warnings,
    }


def qt_document(snapshot: dict[str, Any]) -> QTextDocument:
    """Create an independent rich-text document with publishing styles."""
    document = QTextDocument()
    document.setDefaultStyleSheet(
        "body { font-family: 'Arial'; font-size: 11pt; } "
        "h1,h2,h3,h4,h5,h6 { page-break-after: avoid; } "
        "table { border-collapse: collapse; } td,th { padding: 5px; }"
    )
    document.setHtml(snapshot["html"])
    document.setMetaInformation(
        QTextDocument.MetaInformation.DocumentTitle, snapshot["title"]
    )
    return document


def write_pdf(snapshot: dict[str, Any], path: Path) -> None:
    """Render a paginated PDF without an external office installation."""
    writer = QPdfWriter(str(path))
    writer.setTitle(snapshot["title"])
    writer.setCreator("ProjektKraken")
    writer.setResolution(96)
    size = (
        QPageSize.PageSizeId.Letter
        if snapshot["page_size"] == "Letter"
        else QPageSize.PageSizeId.A4
    )
    writer.setPageLayout(
        QPageLayout(
            QPageSize(size), QPageLayout.Orientation.Portrait, QMarginsF(20, 20, 20, 20)
        )
    )
    document = qt_document(snapshot)
    document.print_(writer)


def write_docx(snapshot: dict[str, Any], path: Path) -> None:
    """Render the same document tree with native Word paragraphs, tables and images."""
    from docx import Document
    from docx.shared import Mm, Pt

    document = Document()
    document.core_properties.title = snapshot["title"]
    section = document.sections[0]
    section.page_width = Mm(216 if snapshot["page_size"] == "Letter" else 210)
    section.page_height = Mm(279 if snapshot["page_size"] == "Letter" else 297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(20)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    parser = _TreeParser()
    parser.feed(snapshot["html"])
    renderer = _WordRenderer()
    for node in parser.root:
        renderer.block(node, document)
    document.save(str(path))


class _WordRenderer:
    """Render document nodes with consistent inline styling and internal links."""

    def __init__(self) -> None:
        """Initialize the per-document bookmark counter."""
        self.bookmark_index = 0

    def inline(
        self, node: Element, paragraph: Any, bold: bool = False, italic: bool = False
    ) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches

        if node.get("id"):
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(self.bookmark_index))
            start.set(qn("w:name"), node.get("id", "").replace("-", "_"))
            paragraph._p.append(start)
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(self.bookmark_index))
            paragraph._p.append(end)
            self.bookmark_index += 1
        if node.tag == "img":
            image_path = QUrl(node.get("src", "")).toLocalFile()
            paragraph.add_run().add_picture(image_path, width=Inches(5.9))
            return
        if node.tag == "br":
            paragraph.add_run().add_break()
        if node.tag == "a" and node.get("href", "").startswith("#"):
            link = OxmlElement("w:hyperlink")
            link.set(qn("w:anchor"), node.get("href", "")[1:].replace("-", "_"))
            run = OxmlElement("w:r")
            value = OxmlElement("w:t")
            value.text = "".join(node.itertext())
            run.append(value)
            link.append(run)
            paragraph._p.append(link)
            return
        if node.text:
            run = paragraph.add_run(node.text)
            run.bold, run.italic = bold, italic
            if node.tag == "code":
                run.font.name = "Consolas"
        for child in node:
            self.inline(
                child,
                paragraph,
                bold or child.tag in {"strong", "b"},
                italic or child.tag in {"em", "i"},
            )
            if child.tail:
                run = paragraph.add_run(child.tail)
                run.bold, run.italic = bold, italic

    def block(self, node: Element, container: Any) -> None:
        if node.tag == "table":
            rows = list(node.iter("tr"))
            if rows:
                table = container.add_table(
                    rows=len(rows), cols=max(len(row) for row in rows)
                )
                table.style = "Table Grid"
                for i, row in enumerate(rows):
                    for j, cell in enumerate(row):
                        self.inline(
                            cell,
                            table.cell(i, j).paragraphs[0],
                            cell.tag == "th",
                        )
        elif node.tag in {"ul", "ol"}:
            for child in node:
                paragraph = container.add_paragraph(
                    style="List Bullet" if node.tag == "ul" else "List Number"
                )
                self.inline(child, paragraph)
        else:
            heading = re.fullmatch(r"h([1-6])", node.tag)
            paragraph = (
                container.add_heading(level=int(heading[1]))
                if heading
                else container.add_paragraph()
            )
            self.inline(node, paragraph)
